import csv
import os
from docx import Document

def reemplazar_texto_en_documento(doc, texto_viejo, texto_nuevo):
    """
    Busca un texto en los párrafos y tablas del documento y lo reemplaza,
    manteniendo el formato original (como el subrayado amarillo).
    """
    # 1. Buscar en los párrafos normales
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if texto_viejo in run.text:
                run.text = run.text.replace(texto_viejo, texto_nuevo)
                
    # 2. Buscar dentro de las tablas (si tu documento tiene tablas)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if texto_viejo in run.text:
                            run.text = run.text.replace(texto_viejo, texto_nuevo)

def main():
    archivo_word = 'plantilla.docx'  # Cambia esto por el nombre de tu Word
    archivo_csv = 'nombres.csv'      # Cambia esto por el nombre de tu CSV
    palabra_clave = 'nombre'         # La palabra exacta que quieres reemplazar
    
    # Crear una carpeta para guardar los archivos generados y no desordenar todo
    carpeta_salida = 'Documentos_Generados'
    os.makedirs(carpeta_salida, exist_ok=True)

    # Leer el archivo CSV
    with open(archivo_csv, mode='r', encoding='utf-8') as file:
        lector_csv = csv.reader(file)
        
        # Iterar sobre cada fila del CSV
        for fila in lector_csv:
            if not fila:
                continue  # Saltar líneas vacías
                
            nombre_persona = fila[0].strip() # Tomar el nombre y limpiar espacios extra
            
            # Cargar el documento plantilla de Word original
            doc = Document(archivo_word)
            
            # Reemplazar la palabra
            reemplazar_texto_en_documento(doc, palabra_clave, nombre_persona)
            
            # Guardar el documento nuevo con el nombre de la persona
            ruta_guardado = os.path.join(carpeta_salida, f"CONVENIO MODIFICATORIO DÍA DE PAGO 2026 GSM {nombre_persona}.docx")
            doc.save(ruta_guardado)
            
            print(f"✅ Archivo creado exitosamente: {ruta_guardado}")

if __name__ == "__main__":
    main()